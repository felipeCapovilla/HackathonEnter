"""
Converte a planilha da base histórica nos CSVs usados pelo gerador da tabela,
pelo backtest e pela prévia de contrato do admin.

Uso:
    python -m scripts.exportar_base_csv --xlsx caminho/Hackaton_Enter_Base_Candidatos.xlsx
"""
from __future__ import annotations

import argparse
import csv
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--xlsx", type=Path, required=True)
    parser.add_argument("--saida", type=Path, default=ROOT / "data")
    args = parser.parse_args()
    args.saida.mkdir(parents=True, exist_ok=True)

    planilha = openpyxl.load_workbook(args.xlsx, read_only=True, data_only=True)
    with open(args.saida / "resultados.csv", "w", newline="", encoding="utf-8") as arquivo:
        escritor = csv.writer(arquivo)
        for linha in planilha["Resultados dos processos"].iter_rows(values_only=True):
            escritor.writerow(linha)
    linhas = list(planilha["Subsídios disponibilizados"].iter_rows(values_only=True))
    with open(args.saida / "subsidios.csv", "w", newline="", encoding="utf-8") as arquivo:
        escritor = csv.writer(arquivo)
        for linha in linhas[1:]:  # a primeira linha da aba é uma legenda
            escritor.writerow(linha)
    print(f"CSVs gravados em {args.saida}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
